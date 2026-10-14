
import csv
import logging
from collections import defaultdict
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from database import get_price_history

PROJECT_DIR = Path(__file__).parent
CSV_PATH = PROJECT_DIR / "price_report.csv"
EXCEL_PATH = PROJECT_DIR / "price_report.xlsx"

COLUMNS = [
    "product_name",
    "currency",
    "previous_price",
    "latest_price",
    "price_change",
    "status",
    "previous_collected_at",
    "latest_collected_at",
]


def build_report_rows():
    """Build a price comparison report from the stored history."""
    history = get_price_history()
    grouped = defaultdict(list)

    for record in history:
        grouped[record["product_name"]].append(record)

    rows = []

    for product_name in sorted(grouped):
        records = grouped[product_name]

        latest = records[0]
        previous = records[1] if len(records) > 1 else None

        latest_price = Decimal(latest["price"])
        previous_price = (
            Decimal(previous["price"]) if previous else None
        )

        if previous_price is None:
            price_change = None
            status = "NO PREVIOUS HISTORY"
        else:
            price_change = latest_price - previous_price

            if price_change > 0:
                status = "PRICE INCREASED"
            elif price_change < 0:
                status = "PRICE DECREASED"
            else:
                status = "NO CHANGE"

        rows.append({
            "product_name": product_name,
            "currency": latest["currency"],
            "previous_price": (
                str(previous_price) if previous_price is not None else ""
            ),
            "latest_price": str(latest_price),
            "price_change": (
                str(price_change) if price_change is not None else ""
            ),
            "status": status,
            "previous_collected_at": (
                previous["collected_at"] if previous else ""
            ),
            "latest_collected_at": latest["collected_at"],
        })

    return rows


def export_csv(rows):
    """Export the report as a CSV file."""
    with CSV_PATH.open("w", newline="", encoding="utf-8-sig") as file:
        writer = csv.DictWriter(file, fieldnames=COLUMNS)
        writer.writeheader()
        writer.writerows(rows)

    logging.info("CSV report exported to %s", CSV_PATH)


def export_excel(rows):
    """Export the report as a formatted Excel workbook."""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Price Report"

    worksheet.append(COLUMNS)

    for row in rows:
        worksheet.append([row[column] for column in COLUMNS])

    # Style the header row.
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78",
        )
        cell.alignment = Alignment(horizontal="center")

    # Format prices and price changes as numbers.
    for row_number in range(2, worksheet.max_row + 1):
        for column_number in (3, 4, 5):
            cell = worksheet.cell(
                row=row_number,
                column=column_number,
            )

            if cell.value not in (None, ""):
                cell.value = float(Decimal(str(cell.value)))
                cell.number_format = "0.00"

    # Make columns readable.
    for column_number, column_name in enumerate(COLUMNS, start=1):
        max_length = len(column_name)

        for cell in worksheet[get_column_letter(column_number)]:
            if cell.value is not None:
                max_length = max(max_length, len(str(cell.value)))

        worksheet.column_dimensions[
            get_column_letter(column_number)
        ].width = min(max_length + 2, 40)

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    workbook.save(EXCEL_PATH)
    logging.info("Excel report exported to %s", EXCEL_PATH)


def export_price_report():
    """Export both CSV and Excel price reports."""
    rows = build_report_rows()

    export_csv(rows)
    export_excel(rows)

    print(f"Exported {len(rows)} products to:")
    print(CSV_PATH)
    print(EXCEL_PATH)


if __name__ == "__main__":
    export_price_report()